from pathlib import Path

import httpx
from docx import Document
from pypdf import PdfWriter

from app.services.source_acquisition import (
    check_url_policy,
    extract_pdf_source,
    extract_url_source,
    extract_word_source,
    ingest_manual_text,
)


def test_p3_s1_manual_source_intake_preserves_stable_id_text_and_span():
    text = "私たちは2018年に創業しました。地域に根ざした支援を続けています。" * 4

    first = ingest_manual_text(text, "会社メモ")
    second = ingest_manual_text(text, "会社メモ")

    assert first.source_id == second.source_id
    assert first.source_type == "manual"
    assert first.source_spans[0].location == "manual:1"
    assert first.metadata["source_priority"] == 1
    assert first.can_proceed is True


def test_p3_s2_url_public_html_extraction_removes_boilerplate_and_records_metadata():
    html = """
    <html>
      <head>
        <title>会社紹介</title>
        <meta property="article:published_time" content="2026-05-01">
      </head>
      <body>
        <nav>不要なナビ</nav>
        <article>
          <h1>会社紹介</h1>
          <p>私たちは地域の相談を受けています。サービスの利用前に状況を確認します。</p>
          <p>担当者が内容を整理し、必要な情報を説明します。公開ページに書かれた範囲だけを材料にします。</p>
          <p>初回の問い合わせでは、利用目的、現在の課題、希望する進め方を確認します。</p>
          <p>確認した内容は担当者が整理し、次に必要な手続きや準備物を案内します。</p>
          <p>本文は公開 HTML の article 要素から抽出され、ナビゲーションやフッターは含めません。</p>
        </article>
        <footer>不要なフッター</footer>
      </body>
    </html>
    """

    transport = httpx.MockTransport(lambda request: httpx.Response(200, text=html, request=request))
    with httpx.Client(transport=transport) as client:
        result = extract_url_source("https://example.com/blog/company", client=client)

    assert result.source_type == "url"
    assert result.title == "会社紹介"
    assert "不要なナビ" not in result.extracted_text
    assert "不要なフッター" not in result.extracted_text
    assert result.metadata["fetch_status"] == 200
    assert result.metadata["published_or_updated_at"] == "2026-05-01"
    assert result.extraction_confidence in {"medium", "high"}
    assert result.can_proceed is True


def test_url_extraction_prefers_richer_main_when_article_is_thin():
    html = """
    <html>
      <head><title>サービス紹介</title></head>
      <body>
        <header>ヘッダー情報</header>
        <main>
          <h1>データ入力スキャニング</h1>
          <article><p>前処理から後処理まで一貫対応</p></article>
          <section>
            <h2>このようなお悩みに対応します</h2>
            <p>データ活用の目的はあるがやり方がわからない担当者を支援します。</p>
            <h2>導入の流れ</h2>
            <p>お問い合わせ、ヒアリング、お見積り、発注、入力、納品まで確認します。</p>
            <p>本文量が article 単体より多い main 要素を選ぶ必要があります。</p>
          </section>
        </main>
      </body>
    </html>
    """
    transport = httpx.MockTransport(lambda request: httpx.Response(200, text=html, request=request))

    with httpx.Client(transport=transport) as client:
        result = extract_url_source("https://example.com/service", client=client)

    assert "導入の流れ" in result.extracted_text
    assert "article 単体より多い main 要素" in result.extracted_text
    assert "ヘッダー情報" not in result.extracted_text


def test_p3_s3_word_minimal_extraction_preserves_blocks_and_tables(tmp_path: Path):
    path = tmp_path / "service.docx"
    doc = Document()
    doc.add_heading("サービス概要", level=1)
    doc.add_paragraph("私たちは初回相談で状況を確認します。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "対象"
    table.cell(0, 1).text = "法人"
    doc.save(path)

    result = extract_word_source(path)

    assert result.source_type == "word"
    assert any(span.location == "block:1" for span in result.source_spans)
    assert any(span.location == "table:1" for span in result.source_spans)
    assert "サービス概要" in result.extracted_text
    assert "対象 | 法人" in result.extracted_text


def test_p3_s3_pdf_placeholder_warns_when_text_is_not_extractable(tmp_path: Path):
    path = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with path.open("wb") as handle:
        writer.write(handle)

    result = extract_pdf_source(path)

    assert result.source_type == "pdf"
    assert result.extracted_text == ""
    assert result.extraction_confidence == "low"
    assert result.can_proceed is False
    assert any("no readable text" in warning for warning in result.warnings)


def test_p3_s4_source_confidence_and_policy_warnings_block_restricted_or_thin_url():
    allowed, warnings = check_url_policy("https://note.com/api/v2/articles/abc")
    thin_html = "<html><body><article><p>短い本文</p></article></body></html>"
    transport = httpx.MockTransport(lambda request: httpx.Response(200, text=thin_html, request=request))

    with httpx.Client(transport=transport) as client:
        result = extract_url_source("https://example.com/thin", client=client)

    assert allowed is False
    assert any("/api/" in warning for warning in warnings)
    assert result.extraction_confidence == "low"
    assert result.can_proceed is False
    assert any("low" in warning for warning in result.warnings)
