from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Dict, Iterable, List

from core.application.markdown_report_service import (
    EXPORTS_DIR,
    build_detailed_markdown_report,
    _as_dict,
    _first_present,
)
from core.application.time_display import current_jst_filename_timestamp


_URL_PATTERN = re.compile(r"https?://[^\s`|)]+")


def _require_docx():
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.opc.constants import RELATIONSHIP_TYPE
        from docx.shared import Inches, Pt, RGBColor
    except Exception as exc:  # pragma: no cover - exercised when dependency is missing.
        raise RuntimeError("DOCX出力には python-docx が必要です。requirements.txt を更新してインストールしてください。") from exc
    return {
        "Document": Document,
        "WD_SECTION": WD_SECTION,
        "WD_TABLE_ALIGNMENT": WD_TABLE_ALIGNMENT,
        "WD_CELL_VERTICAL_ALIGNMENT": WD_CELL_VERTICAL_ALIGNMENT,
        "WD_ALIGN_PARAGRAPH": WD_ALIGN_PARAGRAPH,
        "OxmlElement": OxmlElement,
        "qn": qn,
        "RELATIONSHIP_TYPE": RELATIONSHIP_TYPE,
        "Inches": Inches,
        "Pt": Pt,
        "RGBColor": RGBColor,
    }


def _clean_inline(text: str) -> str:
    return (
        text.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("\\|", "|")
        .replace("`", "")
        .strip()
    )


def _set_run_font(run: Any, *, name: str = "Yu Gothic", size: Any | None = None, bold: bool | None = None, color: Any | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    try:
        run._element.rPr.rFonts.set(run._parent.part.element.nsmap.get("w", "w") + ":eastAsia", name)  # type: ignore[attr-defined]
    except Exception:
        try:
            from docx.oxml.ns import qn

            run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        except Exception:
            pass


def _style_font(style: Any, *, name: str, size: Any, color: Any | None = None, bold: bool | None = None) -> None:
    style.font.name = name
    style.font.size = size
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    try:
        style.element.rPr.rFonts.set(style.element.nsmap.get("w", "w") + ":eastAsia", name)
    except Exception:
        try:
            from docx.oxml.ns import qn

            style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
        except Exception:
            pass


def _configure_document(doc: Any, dep: Dict[str, Any]) -> None:
    Inches = dep["Inches"]
    Pt = dep["Pt"]
    RGBColor = dep["RGBColor"]

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    _style_font(styles["Normal"], name="Yu Gothic", size=Pt(10.5), color=RGBColor(47, 36, 29))
    styles["Normal"].paragraph_format.space_after = Pt(5)
    styles["Normal"].paragraph_format.line_spacing = 1.22

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        _style_font(style, name="Yu Gothic", size=Pt(size), color=RGBColor.from_string(color), bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def _add_hyperlink(paragraph: Any, text: str, url: str, dep: Dict[str, Any]) -> None:
    OxmlElement = dep["OxmlElement"]
    qn = dep["qn"]
    RELATIONSHIP_TYPE = dep["RELATIONSHIP_TYPE"]

    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Yu Gothic")
    fonts.set(qn("w:hAnsi"), "Yu Gothic")
    fonts.set(qn("w:eastAsia"), "Yu Gothic")
    r_pr.append(fonts)
    new_run.append(r_pr)

    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _append_text_with_links(paragraph: Any, text: str, dep: Dict[str, Any], *, bold: bool = False) -> None:
    Pt = dep["Pt"]
    RGBColor = dep["RGBColor"]

    cursor = 0
    for match in _URL_PATTERN.finditer(text):
        before = text[cursor : match.start()]
        if before:
            run = paragraph.add_run(before)
            _set_run_font(run, size=Pt(10.5), bold=bold, color=RGBColor(47, 36, 29))
        url = match.group(0)
        _add_hyperlink(paragraph, url, url, dep)
        cursor = match.end()
    after = text[cursor:]
    if after:
        run = paragraph.add_run(after)
        _set_run_font(run, size=Pt(10.5), bold=bold, color=RGBColor(47, 36, 29))


def _split_markdown_row(line: str) -> List[str]:
    body = line.strip().strip("|")
    cells: List[str] = []
    current: List[str] = []
    escaped = False
    for char in body:
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "|":
            cells.append(_clean_inline("".join(current)))
            current = []
        else:
            current.append(char)
    cells.append(_clean_inline("".join(current)))
    return [cell.strip() for cell in cells]


def _is_separator_row(line: str) -> bool:
    stripped = line.strip().strip("|").strip()
    return bool(stripped) and all(part.strip().replace("-", "").replace(":", "") == "" for part in stripped.split("|"))


def _iter_markdown_blocks(markdown: str) -> Iterable[tuple[str, Any]]:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.strip().startswith("|") and index + 1 < len(lines) and _is_separator_row(lines[index + 1]):
            rows = [_split_markdown_row(line)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_markdown_row(lines[index]))
                index += 1
            yield "table", rows
            continue
        yield "line", line
        index += 1


def _set_cell_shading(cell: Any, fill: str, dep: Dict[str, Any]) -> None:
    OxmlElement = dep["OxmlElement"]
    qn = dep["qn"]
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_margins(table: Any, dep: Dict[str, Any]) -> None:
    OxmlElement = dep["OxmlElement"]
    qn = dep["qn"]
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def _add_table(doc: Any, rows: List[List[str]], dep: Dict[str, Any]) -> None:
    if not rows:
        return

    Inches = dep["Inches"]
    Pt = dep["Pt"]
    RGBColor = dep["RGBColor"]
    WD_TABLE_ALIGNMENT = dep["WD_TABLE_ALIGNMENT"]
    WD_CELL_VERTICAL_ALIGNMENT = dep["WD_CELL_VERTICAL_ALIGNMENT"]

    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_cell_margins(table, dep)

    if column_count <= 2:
        widths = [Inches(1.55), Inches(5.35)]
    elif column_count == 6:
        widths = [Inches(0.35), Inches(0.78), Inches(1.3), Inches(1.75), Inches(1.7), Inches(0.72)]
    else:
        widths = [Inches(6.6 / column_count)] * column_count

    for row_index, row in enumerate(rows):
        for col_index in range(column_count):
            cell = table.cell(row_index, col_index)
            cell.width = widths[min(col_index, len(widths) - 1)]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _set_cell_shading(cell, "E8EEF5", dep)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            text = row[col_index] if col_index < len(row) else ""
            _append_text_with_links(paragraph, text, dep, bold=row_index == 0)
            for run in paragraph.runs:
                _set_run_font(
                    run,
                    size=Pt(8.2 if column_count >= 5 else 9.5),
                    bold=row_index == 0,
                    color=RGBColor(31, 54, 79) if row_index == 0 else RGBColor(47, 36, 29),
                )

    doc.add_paragraph()


def _add_line(doc: Any, line: str, dep: Dict[str, Any]) -> None:
    Pt = dep["Pt"]
    RGBColor = dep["RGBColor"]

    stripped = line.strip()
    if not stripped:
        return
    if stripped == "---":
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run("―" * 24)
        _set_run_font(run, size=Pt(8), color=RGBColor(155, 122, 96))
        return

    if stripped.startswith("# "):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(_clean_inline(stripped[2:]))
        _set_run_font(run, size=Pt(20), bold=True, color=RGBColor(47, 36, 29))
        return
    if stripped.startswith("## "):
        doc.add_heading(_clean_inline(stripped[3:]), level=1)
        return
    if stripped.startswith("### "):
        doc.add_heading(_clean_inline(stripped[4:]), level=2)
        return
    if stripped.startswith("#### "):
        doc.add_heading(_clean_inline(stripped[5:]), level=3)
        return

    bullet_level = 0
    bullet_text = ""
    if stripped.startswith("- "):
        bullet_text = stripped[2:]
    elif line.startswith("  - "):
        bullet_level = 1
        bullet_text = line.strip()[2:]

    if bullet_text:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = dep["Inches"](0.35 + 0.18 * bullet_level)
        paragraph.paragraph_format.first_line_indent = dep["Inches"](-0.16)
        paragraph.paragraph_format.space_after = Pt(3)
        _append_text_with_links(paragraph, _clean_inline(bullet_text), dep)
        return

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    _append_text_with_links(paragraph, _clean_inline(stripped), dep)


def build_detailed_docx_report(bundle: Dict[str, Any]) -> Any:
    dep = _require_docx()
    doc = dep["Document"]()
    _configure_document(doc, dep)

    markdown = build_detailed_markdown_report(bundle)
    for block_type, payload in _iter_markdown_blocks(markdown):
        if block_type == "table":
            _add_table(doc, payload, dep)
        else:
            _add_line(doc, payload, dep)
    return doc


def export_detailed_docx_report(bundle: Dict[str, Any]) -> Path:
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    run_row = _as_dict(bundle.get("run"))
    snapshot = _as_dict(bundle.get("snapshot"))
    meta = _as_dict(snapshot.get("meta"))
    run_id = _first_present(run_row.get("id"), meta.get("run_id"), "unknown")
    timestamp = current_jst_filename_timestamp()
    output_path = EXPORTS_DIR / f"detailed-report-run-{run_id}-{timestamp}.docx"
    doc = build_detailed_docx_report(bundle)
    doc.save(output_path)
    return output_path
