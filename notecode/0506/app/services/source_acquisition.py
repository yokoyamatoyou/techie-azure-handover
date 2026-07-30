from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import httpx
import yaml
from bs4 import BeautifulSoup
from docx import Document

from app.services.pdf_text_extractor import extract_pdf_text


ROOT = Path(__file__).resolve().parents[2]
CONFIG_PATH = ROOT / "app" / "config" / "source_acquisition.yaml"


@dataclass(frozen=True)
class SourceSpan:
    span_id: str
    text: str
    location: str


@dataclass(frozen=True)
class ExtractedSource:
    source_id: str
    source_type: str
    title: str
    extracted_text: str
    source_spans: list[SourceSpan]
    metadata: dict[str, Any]
    warnings: list[str] = field(default_factory=list)
    extraction_confidence: str = "low"
    can_proceed: bool = False


def load_source_config(path: Path = CONFIG_PATH) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        data = yaml.safe_load(handle) or {}
    if not isinstance(data, dict):
        raise ValueError("source acquisition config must be a mapping")
    return data


def stable_source_id(source_type: str, title: str, text_or_locator: str) -> str:
    digest = hashlib.sha256(f"{source_type}|{title}|{text_or_locator}".encode("utf-8")).hexdigest()[:12]
    return f"{source_type}_{digest}"


def _confidence_for_text(text: str, config: dict[str, Any]) -> str:
    length = len(text.strip())
    thresholds = config.get("confidence", {})
    if length >= int(thresholds.get("high_min_chars", 500)):
        return "high"
    if length >= int(thresholds.get("medium_min_chars", 120)):
        return "medium"
    return "low"


def _can_proceed(confidence: str) -> bool:
    return confidence in {"high", "medium"}


def ingest_manual_text(text: str, title: str = "manual source") -> ExtractedSource:
    config = load_source_config()
    clean_text = text.strip()
    confidence = _confidence_for_text(clean_text, config)
    warnings = [] if clean_text else ["manual text is empty"]
    if confidence == "low" and clean_text:
        warnings.append("manual text is thin and should be reviewed before generation")
    return ExtractedSource(
        source_id=stable_source_id("manual", title, clean_text),
        source_type="manual",
        title=title,
        extracted_text=clean_text,
        source_spans=[SourceSpan("manual_001", clean_text, "manual:1")] if clean_text else [],
        metadata={
            "source_label": title,
            "source_priority": config["source_priority"]["manual"],
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "extraction_method": "manual_text",
        },
        warnings=warnings,
        extraction_confidence=confidence,
        can_proceed=_can_proceed(confidence),
    )


def check_url_policy(url: str, config: dict[str, Any] | None = None) -> tuple[bool, list[str]]:
    cfg = config or load_source_config()
    parsed = urlparse(url)
    warnings: list[str] = []
    if parsed.scheme not in set(cfg.get("url_policy", {}).get("allowed_schemes", [])):
        return False, [f"blocked URL scheme: {parsed.scheme}"]

    lower_path = parsed.path.lower()
    for fragment in cfg.get("url_policy", {}).get("blocked_path_fragments", []):
        if fragment in lower_path:
            warnings.append(f"blocked path fragment: {fragment}")

    if parsed.netloc.endswith("note.com") and "/api/" in lower_path:
        warnings.append("note internal API endpoints are not allowed")
    if "hatena" in parsed.netloc and any(part in lower_path for part in ("/admin", "/login")):
        warnings.append("Hatena authenticated/admin pages are not allowed")

    return not warnings, warnings


def extract_url_source(url: str, client: httpx.Client | None = None) -> ExtractedSource:
    config = load_source_config()
    allowed, policy_warnings = check_url_policy(url, config)
    if not allowed:
        return ExtractedSource(
            source_id=stable_source_id("url", url, url),
            source_type="url",
            title=url,
            extracted_text="",
            source_spans=[],
            metadata={"url": url, "extraction_method": "public_html"},
            warnings=policy_warnings,
            extraction_confidence="low",
            can_proceed=False,
        )

    owns_client = client is None
    active_client = client or httpx.Client(follow_redirects=True, timeout=10.0)
    try:
        response = active_client.get(url)
        response.raise_for_status()
    finally:
        if owns_client:
            active_client.close()

    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "aside", "form", "header"]):
        tag.decompose()

    title = _extract_title(soup) or url
    published = _extract_published_at(soup)
    container = _select_content_container(soup)
    text = container.get_text("\n", strip=True)
    confidence = _confidence_for_text(text, config)
    warnings = policy_warnings.copy()
    if confidence == "low":
        warnings.append("URL extraction confidence is low; do not proceed silently")

    return ExtractedSource(
        source_id=stable_source_id("url", title, url),
        source_type="url",
        title=title,
        extracted_text=text,
        source_spans=[SourceSpan("url_001", text, "html:article")] if text else [],
        metadata={
            "url": url,
            "canonical_url": str(response.url),
            "published_or_updated_at": published,
            "source_priority": config["source_priority"]["external_url"],
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "fetch_status": response.status_code,
            "extraction_method": "public_html",
        },
        warnings=warnings,
        extraction_confidence=confidence,
        can_proceed=_can_proceed(confidence),
    )


def _extract_title(soup: BeautifulSoup) -> str | None:
    og_title = soup.find("meta", property="og:title")
    if og_title and og_title.get("content"):
        return str(og_title["content"]).strip()
    if soup.title and soup.title.string:
        return soup.title.string.strip()
    heading = soup.find("h1")
    return heading.get_text(" ", strip=True) if heading else None


def _select_content_container(soup: BeautifulSoup) -> Any:
    candidates = []
    for selector in ("main", "article", "[role='main']", ".main", ".contents", ".content"):
        candidates.extend(soup.select(selector))
    if soup.body:
        candidates.append(soup.body)
    candidates.append(soup)

    def score(node: Any) -> int:
        text = node.get_text("\n", strip=True)
        heading_bonus = len(node.find_all(["h1", "h2", "h3", "dt", "dd"])) * 20 if hasattr(node, "find_all") else 0
        return len(text) + heading_bonus

    return max(candidates, key=score)


def _extract_published_at(soup: BeautifulSoup) -> str | None:
    time_tag = soup.find("time")
    if time_tag and time_tag.get("datetime"):
        return str(time_tag["datetime"])
    for attr in ("article:published_time", "article:modified_time"):
        meta = soup.find("meta", property=attr)
        if meta and meta.get("content"):
            return str(meta["content"])
    return None


def extract_pdf_source(path: Path) -> ExtractedSource:
    config = load_source_config()
    extraction = extract_pdf_text(path)
    spans = [
        SourceSpan(span_id=span_id, text=text, location=location)
        for span_id, text, location in extraction.spans
    ]
    text = "\n\n".join(span.text for span in spans)
    confidence = _confidence_for_text(text, config)
    warnings = list(extraction.warnings)
    if not text:
        warnings.append("PDF text extraction produced no readable text")
    if confidence == "low":
        warnings.append("PDF extraction confidence is low; review source before generation")
    title = path.stem
    return ExtractedSource(
        source_id=stable_source_id("pdf", title, str(path.resolve())),
        source_type="pdf",
        title=title,
        extracted_text=text,
        source_spans=spans,
        metadata={
            "source_label": path.name,
            "source_priority": config["source_priority"]["pdf"],
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            **extraction.metadata,
        },
        warnings=warnings,
        extraction_confidence=confidence,
        can_proceed=_can_proceed(confidence),
    )


def extract_word_source(path: Path) -> ExtractedSource:
    config = load_source_config()
    document = Document(str(path))
    spans: list[SourceSpan] = []
    block_index = 1
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            spans.append(SourceSpan(f"word_{block_index:03d}", text, f"block:{block_index}"))
            block_index += 1
    for table_index, table in enumerate(document.tables, start=1):
        cells = [cell.text.strip() for row in table.rows for cell in row.cells if cell.text.strip()]
        if cells:
            spans.append(SourceSpan(f"word_table_{table_index:03d}", " | ".join(cells), f"table:{table_index}"))

    text = "\n".join(span.text for span in spans)
    confidence = _confidence_for_text(text, config)
    warnings: list[str] = []
    if confidence == "low":
        warnings.append("Word extraction confidence is low; review source before generation")
    title = path.stem
    return ExtractedSource(
        source_id=stable_source_id("word", title, str(path.resolve())),
        source_type="word",
        title=title,
        extracted_text=text,
        source_spans=spans,
        metadata={
            "source_label": path.name,
            "source_priority": config["source_priority"]["word"],
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "extraction_method": "word_text",
        },
        warnings=warnings,
        extraction_confidence=confidence,
        can_proceed=_can_proceed(confidence),
    )
